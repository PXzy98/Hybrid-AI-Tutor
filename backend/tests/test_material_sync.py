"""Uploaded course text is searchable, refreshable, persistent and accurately listed."""

from concurrent.futures import ThreadPoolExecutor

import pytest

from app.rag.documents import DocumentError, extract_text
from app.rag.medical_hybrid import MedicalHybridRetriever

TOPIC = "asthma"
TEXT = "Quartzzyme is the fictional keyword used only in the uploaded course lecture on airway assessment."


def upload_path(root, name="lecture.txt", text=TEXT):
    p = root / TOPIC / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(text, encoding="utf-8")
    return p


def uploaded_rows(r):
    return [row for row in r.material_status(TOPIC) if row["origin"] == "uploaded"]


def test_builtin_inventory_is_present_for_each_topic(tmp_path):
    r = MedicalHybridRetriever(tmp_path)
    assert len(r.indices) == 5
    for topic in r.indices:
        rows = r.material_status(topic)
        assert len(rows) == 1 and rows[0]["origin"] == "builtin" and rows[0]["indexed"]
        assert rows[0]["passage_count"] > 0 and rows[0]["source"].startswith("https://medlineplus.gov/")


@pytest.mark.parametrize("extension", [".txt", ".md"])
def test_upload_is_retrievable_and_survives_restart(tmp_path, extension):
    r = MedicalHybridRetriever(tmp_path)
    path = upload_path(tmp_path, "lecture" + extension)
    r.ingest(TOPIC, str(path))
    result = r.retrieve("quartzzyme course lecture", TOPIC)
    assert any(p.source == f"course/{TOPIC}/{path.name}" and "Quartzzyme" in p.text for p in result)
    assert uploaded_rows(r)[0]["indexed"]
    restarted = MedicalHybridRetriever(tmp_path)
    assert any("Quartzzyme" in p.text for p in restarted.retrieve("quartzzyme", TOPIC))
    assert all("Quartzzyme" not in p.text for p in r.retrieve("quartzzyme", "migraine"))


def test_replacement_removes_old_text_and_does_not_duplicate(tmp_path):
    r = MedicalHybridRetriever(tmp_path)
    path = upload_path(tmp_path)
    r.ingest(TOPIC, str(path))
    path.write_text(TEXT.replace("Quartzzyme", "Zephyrlens"), encoding="utf-8")
    r.ingest(TOPIC, str(path))
    r.ingest(TOPIC, str(path))
    texts = " ".join(p.text for p in r.indices[TOPIC].passages)
    assert "Quartzzyme" not in texts and "Zephyrlens" in texts
    assert len(uploaded_rows(r)) == 1


def test_other_worker_and_external_deletion_refresh_without_restart(tmp_path):
    a, b = MedicalHybridRetriever(tmp_path), MedicalHybridRetriever(tmp_path)
    path = upload_path(tmp_path)
    a.ingest(TOPIC, str(path))
    assert any("Quartzzyme" in p.text for p in b.retrieve("quartzzyme", TOPIC))
    path.unlink()
    b.retrieve("quartzzyme", TOPIC)
    assert not uploaded_rows(b)
    assert all("Quartzzyme" not in p.text for p in b.indices[TOPIC].passages)


def test_invalid_replacement_is_not_served_as_stale_content(tmp_path):
    r = MedicalHybridRetriever(tmp_path)
    path = upload_path(tmp_path)
    r.ingest(TOPIC, str(path))
    path.write_bytes(b"\0binary replacement")
    with pytest.raises(DocumentError):
        r.ingest(TOPIC, str(path))
    assert not uploaded_rows(r)[0]["indexed"]
    assert uploaded_rows(r)[0]["message"]
    assert all("Quartzzyme" not in p.text for p in r.indices[TOPIC].passages)
    assert r.retrieve("What is asthma?", TOPIC)


def test_invalid_file_does_not_prevent_startup(tmp_path):
    path = tmp_path / TOPIC / "bad.pdf"
    path.parent.mkdir(parents=True)
    path.write_bytes(b"not a pdf")
    r = MedicalHybridRetriever(tmp_path)
    assert not uploaded_rows(r)[0]["indexed"]
    assert r.retrieve("What is asthma?", TOPIC)


def test_html_does_not_index_scripts_or_fetch_urls(tmp_path):
    path = upload_path(tmp_path, "lecture.html", f"<h1>Lecture</h1><p>{TEXT}</p>"
                       "<script>DO_NOT_INDEX_THIS</script><img src='https://invalid.example/never-fetched'>")
    r = MedicalHybridRetriever(tmp_path)
    assert "DO_NOT_INDEX_THIS" not in " ".join(p.text for p in r.indices[TOPIC].passages)
    assert any("Quartzzyme" in p.text for p in r.retrieve("quartzzyme", TOPIC))
    assert path.exists()


def test_docx_paragraphs_and_tables_are_indexed(tmp_path):
    from docx import Document

    path = tmp_path / TOPIC / "lecture.docx"
    path.parent.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph(TEXT)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Zephyrlens table definition in the course."
    doc.save(path)
    r = MedicalHybridRetriever(tmp_path)
    text = " ".join(p.text for p in r.indices[TOPIC].passages)
    assert "Quartzzyme" in text and "Zephyrlens" in text
    assert uploaded_rows(r)[0]["indexed"]


def test_text_pdf_is_indexed_and_scan_is_explicitly_unindexed(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / TOPIC / "lecture.pdf"
    path.parent.mkdir(parents=True)
    # Small standards-based fixture; no PDF-generation dependency is needed by the test suite.
    stream = f"BT /F1 12 Tf 40 700 Td ({TEXT}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 900 800] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream",
    ]
    data, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(data)
    r = MedicalHybridRetriever(tmp_path)
    assert any("Quartzzyme" in p.text for p in r.retrieve("quartzzyme", TOPIC))
    blank = tmp_path / TOPIC / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with blank.open("wb") as stream:
        writer.write(stream)
    with pytest.raises(DocumentError, match="OCR"):
        r.ingest(TOPIC, str(blank))
    assert any(row["name"] == "scan.pdf" and not row["indexed"] for row in uploaded_rows(r))


def test_utf16_and_short_text_are_not_discarded(tmp_path):
    path = tmp_path / TOPIC / "short.txt"
    path.parent.mkdir(parents=True)
    path.write_text("Quartzzyme means a lecture keyword.", encoding="utf-16")
    r = MedicalHybridRetriever(tmp_path)
    assert uploaded_rows(r)[0]["indexed"]
    assert any("Quartzzyme" in p.text for p in r.retrieve("quartzzyme", TOPIC))


def test_short_final_constraint_survives_chunking(tmp_path):
    upload_path(tmp_path, text=("Airway lecture text " * 32) + "\n" + ("Assessment detail " * 5)
                + "\nKeep Zephyrlens.")
    r = MedicalHybridRetriever(tmp_path)
    text = " ".join(p.text for p in r.indices[TOPIC].passages)
    assert "Keep Zephyrlens." in text


def test_outside_topic_path_is_rejected(tmp_path):
    r = MedicalHybridRetriever(tmp_path / "materials")
    path = tmp_path / "outside.txt"
    path.write_text(TEXT)
    with pytest.raises(DocumentError, match="inside"):
        r.ingest(TOPIC, str(path))


def test_concurrent_readers_and_inventory_get_complete_indexes(tmp_path):
    r = MedicalHybridRetriever(tmp_path)
    path = upload_path(tmp_path)
    with ThreadPoolExecutor(max_workers=4) as pool:
        tasks = [pool.submit(r.ingest, TOPIC, str(path))]
        tasks.extend(pool.submit(r.retrieve, "quartzzyme asthma", TOPIC) for _ in range(8))
        for task in tasks:
            task.result()
    assert uploaded_rows(r)[0]["indexed"]


def test_extraction_budget_is_explicit(tmp_path, monkeypatch):
    monkeypatch.setattr("app.rag.documents.MAX_TEXT_CHARS", 20)
    path = tmp_path / "large.txt"
    path.write_text(TEXT)
    with pytest.raises(DocumentError, match="Split"):
        extract_text(path)


def test_teacher_inventory_upload_and_semantic_answer_use_same_documents(make_client, monkeypatch, tmp_path):
    import numpy as np

    from tests.conftest import Api
    from tests.test_semantic_rag_integration import config, configure_classifier

    configure_classifier(monkeypatch)
    monkeypatch.setattr(
        "app.routing.evidence.pair_scores", lambda name, pairs, cfg: np.full(len(pairs), 0.99)
    )
    settings = config()
    settings["rag"]["materials_dir"] = str(tmp_path / "uploads")
    with make_client(settings) as client:
        api = Api(client)
        api.make_instructor()
        api.make_student()
        api.login("prof")
        content = client.get("/api/v1/admin/content").json()
        assert len(content["topics"]) == 5
        assert all(any(r["origin"] == "builtin" and r["indexed"] for r in t["rag_materials"])
                   for t in content["topics"])
        upload = client.post(f"/api/v1/admin/topics/{TOPIC}/materials", files={
            "file": ("lecture.txt", TEXT.encode(), "text/plain"),
        })
        assert upload.status_code == 201 and upload.json()["ingested"] is True, upload.text
        state = client.get("/api/v1/admin/content").json()
        topic = next(t for t in state["topics"] if t["topic"] == TOPIC)
        row = next(m for m in topic["rag_materials"] if m["origin"] == "uploaded")
        assert row["indexed"] and row["passage_count"] > 0
        api.login("alice")
        assert client.get("/api/v1/admin/content").status_code == 403
        reply = client.post("/api/v1/ask", json={
            "topic": TOPIC, "question": "According to my lecture, what does quartzzyme mean?",
        })
        assert reply.status_code == 200, reply.text
        assert reply.json()["route_target"] == "local"
        assert any(s["source"] == f"course/{TOPIC}/lecture.txt" for s in reply.json()["sources"])
        assert "Quartzzyme" in api.local_llm.sent_text()


def test_api_parse_failure_is_explicit_and_replacement_can_recover(make_client, tmp_path):
    from tests.conftest import Api

    with make_client({"rag": {"impl": "medical_hybrid", "materials_dir": str(tmp_path / "uploads")}}) as c:
        api = Api(c)
        api.make_instructor()
        api.login("prof")
        url = f"/api/v1/admin/topics/{TOPIC}/materials"
        bad = c.post(url, files={"file": ("lecture.txt", b"\0invalid", "text/plain")})
        assert bad.status_code == 201 and bad.json()["ingested"] is False
        assert "binary" in bad.json()["ingest_error"]
        duplicate = c.post(url, files={"file": ("lecture.txt", TEXT.encode(), "text/plain")})
        assert duplicate.status_code == 409
        good = c.post(url + "?replace=true", files={"file": ("lecture.txt", TEXT.encode(), "text/plain")})
        assert good.status_code == 201 and good.json()["ingested"] is True
        assert all(row["message"] is None for row in uploaded_rows(api.services.retriever))
        assert not list((tmp_path / "uploads" / TOPIC).glob(".upload_*"))
