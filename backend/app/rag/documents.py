"""Local text extraction for instructor uploads; never fetch embedded URLs or run OCR."""

import zipfile
from html.parser import HTMLParser
from pathlib import Path

SUFFIXES = {".md", ".txt", ".pdf", ".html", ".htm", ".docx"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 500_000
MAX_PDF_PAGES = 250


class DocumentError(ValueError):
    """Safe, teacher-visible reason a stored document is not searchable."""


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"} and self.hidden:
            self.hidden -= 1
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
    except UnicodeError as error:
        raise DocumentError("Save text or HTML files as UTF-8 or UTF-16 and upload again.") from error


def extract_text(path: Path) -> str:
    if path.suffix.lower() not in SUFFIXES:
        raise DocumentError("This file type is not supported for retrieval.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise DocumentError("The file exceeds the 20 MB limit.")
    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            if reader.is_encrypted:
                raise DocumentError("Remove PDF password protection before uploading.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise DocumentError("Split PDFs into documents of at most 250 pages.")
            parts, count = [], 0
            for page in reader.pages:
                text = page.extract_text() or ""
                count += len(text)
                if count > MAX_TEXT_CHARS:
                    raise DocumentError("Split documents with more than 500,000 extracted characters.")
                parts.append(text)
            text = "\n".join(parts)
        elif suffix == ".docx":
            from docx import Document

            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                    raise DocumentError("The unpacked Word document is too large.")
            document = Document(path)
            parts = [p.text for p in document.paragraphs]
            parts.extend("\t".join(cell.text for cell in row.cells)
                         for table in document.tables for row in table.rows)
            text = "\n".join(parts)
        else:
            text = _decode(path.read_bytes())
            if suffix in {".html", ".htm"}:
                parser = _HTMLText()
                parser.feed(text)
                parser.close()
                text = "".join(parser.parts)
    except DocumentError:
        raise
    except ImportError as error:
        raise DocumentError("Install backend dependencies and restart to read this format.") from error
    except Exception as error:
        raise DocumentError("This document could not be read. Check the file and upload it again.") from error
    if "\0" in text:
        raise DocumentError("The document contains binary data instead of readable text.")
    text = text.strip()
    if not text:
        raise DocumentError("No readable text was found. Scanned PDFs require OCR before uploading.")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentError("Split documents with more than 500,000 extracted characters.")
    return text
